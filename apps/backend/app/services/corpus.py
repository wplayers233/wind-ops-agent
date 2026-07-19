from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from app.services.ocr import build_ocr_bundle
from app.services.vision import build_visual_enrichment

SUPPORTED_SUFFIXES = {".json", ".jsonl", ".txt", ".md", ".pdf", ".pptx", ".png", ".jpg", ".jpeg"}


def chunk_document(document: dict, chunk_size: int = 320) -> list[dict]:
    text = " ".join(str(document.get(key, "")) for key in ("title", "content", "ocr_text", "visual_caption"))
    return [
        {
            "id": f"{document.get('id', 'doc')}-{index}",
            "document_id": document.get("id", ""),
            "page": document.get("page", 1),
            "content": text[index : index + chunk_size],
            "metadata": {key: document.get(key, "") for key in ("model", "component", "fault_domain", "alarm_code", "doc_type", "source")},
        }
        for index in range(0, max(1, len(text)), chunk_size)
    ]


def load_documents(path: str | Path) -> list[dict]:
    source = Path(path)
    if source.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise ValueError(f"unsupported corpus format: {source.suffix}")
    if source.suffix.lower() == ".jsonl":
        return [json.loads(line) for line in source.read_text(encoding="utf-8").splitlines() if line.strip()]
    if source.suffix.lower() == ".json":
        payload = json.loads(source.read_text(encoding="utf-8"))
        return payload if isinstance(payload, list) else [payload]
    if source.suffix.lower() in {".txt", ".md"}:
        return [{"id": source.stem, "title": source.name, "content": source.read_text(encoding="utf-8"), "source": source.name, "page": 1, "content_type": "text"}]
    return parse_binary(source.read_bytes(), source.name, source.suffix.lower())


def parse_binary(data: bytes, filename: str, suffix: str) -> list[dict]:
    if suffix == ".pdf":
        pages: list[dict[str, Any]] = []
        try:
            import fitz

            pdf = fitz.open(stream=data, filetype="pdf")
            for index, page in enumerate(pdf, start=1):
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                pages.append({"id": f"{Path(filename).stem}-{index}", "title": filename, "content": page.get_text("text") or "", "source": filename, "page": index, "content_type": "ocr", "image_bytes": pix.tobytes("png")})
        except (ImportError, RuntimeError):
            from io import BytesIO
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))
            pages = [{"id": f"{Path(filename).stem}-{index}", "title": filename, "content": page.extract_text() or "", "source": filename, "page": index, "content_type": "ocr"} for index, page in enumerate(reader.pages, start=1)]
        return build_visual_enrichment(source=filename, pages=build_ocr_bundle(source=filename, pages=pages))
    if suffix == ".pptx":
        try:
            from io import BytesIO
            from pptx import Presentation

            presentation = Presentation(BytesIO(data))
        except ImportError as exc:
            raise RuntimeError("PPTX parser unavailable; install python-pptx") from exc
        pages = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = []
            regions = []
            image_bytes = None
            for shape_index, shape in enumerate(slide.shapes):
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
                if getattr(shape, "shape_type", None) == 13:
                    try:
                        image_bytes = shape.image.blob
                        regions.append({"id": f"{Path(filename).stem}-p{index}-image-{shape_index}", "modality": "image", "page": index, "source": filename, "bbox": [shape.left, shape.top, shape.width, shape.height], "content_type": "image"})
                    except Exception:
                        pass
                if getattr(shape, "has_table", False):
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    regions.append({"id": f"{Path(filename).stem}-p{index}-table-{shape_index}", "modality": "table", "page": index, "source": filename, "table": rows, "bbox": [shape.left, shape.top, shape.width, shape.height]})
            pages.append({"id": f"{Path(filename).stem}-{index}", "title": filename, "content": " ".join(texts), "source": filename, "page": index, "content_type": "slide", "image_bytes": image_bytes, "regions": regions})
        return build_visual_enrichment(source=filename, pages=build_ocr_bundle(source=filename, pages=pages))
    if suffix in {".png", ".jpg", ".jpeg"}:
        return build_visual_enrichment(source=filename, pages=build_ocr_bundle(source=filename, pages=[{"id": Path(filename).stem, "title": filename, "content": "", "source": filename, "page": 1, "image_bytes": data, "content_type": "image"}]))
    raise ValueError(f"unsupported binary format: {suffix}")


def build_ingestion_report(documents: list[dict], *, provider_mode: str = "auto") -> dict[str, Any]:
    chunks = [chunk for document in documents for chunk in chunk_document(document)]
    regions = sum(len(document.get("regions", [])) for document in documents)
    images = sum(1 for document in documents if document.get("image_bytes") or document.get("content_type") == "image")
    ocr_status = sorted({str(item.get("ocr_status", "offline_fallback")) for item in documents})
    vision_status = sorted({str(item.get("vision_status", "offline_fallback")) for item in documents})
    embedding_status = "auto" if provider_mode == "auto" else "offline_fallback"
    return {
        "document_count": len(documents),
        "page_count": len(documents),
        "chunk_count": len(chunks),
        "region_count": regions,
        "image_count": images,
        "provider_mode": provider_mode,
        "ocr_status": ocr_status,
        "vision_status": vision_status,
        "embedding_status": embedding_status,
        "status": "ready" if documents else "empty",
        "chunks": chunks,
    }
